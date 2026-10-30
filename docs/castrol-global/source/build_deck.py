# Builds ../KAD_Platform_Castrol_Global.pptx (16:9, 21 slides + 1 appendix slide, speaker notes on every slide).
# Values on the slides are illustrative (USD) or demo data only: no real company figures.
# Inputs: ../screenshots/*.png (framed phones, from frames.py), icons/*.png (from assets.mjs), kad-wordmark.png.
# Run: python3 build_deck.py
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from PIL import Image
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "..", "screenshots")
ICONS = os.path.join(HERE, "icons")
OUT = os.path.join(HERE, "..", "KAD_Platform_Castrol_Global.pptx")

GREEN = RGBColor(0x00, 0x93, 0x43)
GREEN_DARK = RGBColor(0x00, 0x6E, 0x32)
GREEN_TINT = RGBColor(0xE6, 0xF4, 0xEC)
GREEN_MID = RGBColor(0x9E, 0xD3, 0xB3)
RED = RGBColor(0xE2, 0x00, 0x1A)
RED_TINT = RGBColor(0xFD, 0xEB, 0xED)
DARK = RGBColor(0x23, 0x27, 0x2B)
GREY = RGBColor(0x5F, 0x66, 0x70)
LIGHT = RGBColor(0xF4, 0xF5, 0xF6)
LINE = RGBColor(0xD9, 0xDE, 0xDB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW = 13.333


def text(slide, x, y, w, h, content, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line=1.0):
    """content: str, or list of paragraphs; a paragraph is a str or a list of (text, {size,bold,color}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if spacing is not None and i > 0:
            p.space_before = Pt(spacing)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, st in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.color.rgb = st.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=None, line=None, radius=0.12, line_w=1.0, dash=False, shape=None):
    shp = slide.shapes.add_shape(shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE), Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and (shape is None):
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def pic(slide, path, x, y, w=None, h=None):
    if w is None and h is None:
        raise ValueError
    im = Image.open(path)
    ar = im.size[0] / im.size[1]
    if w is None:
        w = h * ar
    if h is None:
        h = w / ar
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def phone(slide, name, x, y, h):
    return pic(slide, os.path.join(SHOTS, name + ".png"), x, y, h=h)


def icon(slide, name, x, y, d, bg=GREEN_TINT, color="g", ratio=0.56):
    if bg is not None:
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        c.fill.solid()
        c.fill.fore_color.rgb = bg
        c.line.fill.background()
        c.shadow.inherit = False
    s = d * ratio
    slide.shapes.add_picture(os.path.join(ICONS, f"{name}_{color}.png"), Inches(x + (d - s) / 2), Inches(y + (d - s) / 2), Inches(s), Inches(s))


def stripe(slide, x, y, w=0.9):
    box(slide, x, y, w * 0.72, 0.07, fill=GREEN, radius=0)
    box(slide, x + w * 0.72, y, w * 0.28, 0.07, fill=RED, radius=0)


def header(slide, kicker, headline, size=30):
    stripe(slide, 0.6, 0.38)
    text(slide, 0.6, 0.55, 10, 0.3, kicker.upper(), size=12, bold=True, color=GREEN)
    text(slide, 0.6, 0.85, 12.1, 1.0, headline, size=size, bold=True, color=DARK, line=0.95)


def footer(slide, n):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6), Inches(7.0), Inches(SW - 0.6), Inches(7.0))
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)
    pic(slide, os.path.join(HERE, "kad-wordmark.png"), 0.6, 7.13, h=0.17)
    text(slide, 1.55, 7.1, 7, 0.25, "Digital Sales Platform for Castrol Distribution", size=9, color=GREY)
    text(slide, SW - 1.6, 7.1, 1.0, 0.25, str(n), size=9, color=GREY, align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


def logo_spot(slide, x, y, w, h, on_dark=False):
    col = WHITE if on_dark else GREY
    b = box(slide, x, y, w, h, fill=None, line=col, radius=0.08, line_w=1.25, dash=True)
    text(slide, x, y, w, h, [[("CASTROL LOGO", {"bold": True, "size": 11, "color": col})], [("place official asset here", {"size": 9, "color": col})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return b


def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    return s


_n = [0]


def nxt():
    _n[0] += 1
    return _n[0]


def badge(slide, x, y, label="ILLUSTRATIVE VALUES · NOT COMPANY DATA", w=3.9):
    box(slide, x, y, w, 0.34, fill=RED_TINT, radius=0.17)
    text(slide, x, y, w, 0.34, label, size=10, bold=True, color=RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


DEMO_NOTE = "All screens and example values in this deck are made-up demo data. No real customer, staff or company figures are shown."

# ---------------------------------------------------------------- Title
s = new_slide(); nxt()
box(s, 8.35, 0, SW - 8.35, 7.5, fill=GREEN, radius=0)
box(s, 8.35, 0, 0.09, 7.5, fill=RED, radius=0)
phone(s, "checkin", 7.05, 1.05, 5.55)
phone(s, "vaghinak", 10.55, 1.05, 5.55)
phone(s, "map", 8.75, 0.55, 6.4)
stripe(s, 0.7, 1.2, 1.1)
text(s, 0.7, 1.45, 6.5, 0.35, "PREPARED FOR CASTROL GLOBAL", size=13, bold=True, color=GREEN)
text(s, 0.7, 1.9, 6.6, 2.7, "KAD Motors: Digital Sales Platform for Castrol Distribution", size=40, bold=True, color=DARK, line=0.95)
text(s, 0.7, 4.45, 6.5, 0.9, ["Field app · AI deputy director · CEO bot", "Yerevan, Armenia · October 2026"], size=18, color=GREY, spacing=4)
pic(s, os.path.join(HERE, "kad-wordmark.png"), 0.7, 5.95, h=0.4)
logo_spot(s, 2.75, 5.7, 2.4, 0.9)
notes(s, """
Thank you for your time. KAD Motors is a Castrol distributor in Armenia. Over the last months our own team built a digital platform that runs the whole commercial side of a distributor: a field app for the team, an AI deputy director who controls the daily work, and a CEO bot for management.
Today I will show the real problems of a distributor, how the platform solves each one, and why we believe it can become a model for other Castrol distributors, with shelf and sell-out visibility for Castrol.
""" + DEMO_NOTE)

# ---------------------------------------------------------------- At a glance
s = new_slide(); n = nxt()
header(s, "KAD Motors at a glance", "A Castrol distributor with one goal: No. 1 in Armenia")
facts = [("store", "Castrol distributor in Armenia"), ("team", "Field, warehouse and finance teams"), ("mapPinCheck", "Customers across Armenia's regions")]
for i, (ic, lab) in enumerate(facts):
    x = 0.6 + i * 2.27
    box(s, x, 2.0, 2.1, 2.05, fill=LIGHT)
    icon(s, ic, x + 0.22, 2.22, 0.75)
    text(s, x + 0.22, 3.12, 1.75, 0.9, lab, size=13.5, bold=True, color=DARK)
text(s, 0.6, 4.45, 6.6, 0.3, "WHO WE SERVE", size=12, bold=True, color=GREY)
segs = [("gauge", "Workshops & garages"), ("repeat", "Oil-change points"), ("store", "Auto-parts shops"), ("truck", "Fleets & car dealers")]
for i, (ic, lab) in enumerate(segs):
    x = 0.6 + i * 1.7
    icon(s, ic, x + 0.38, 4.9, 0.78)
    text(s, x, 5.78, 1.55, 0.6, lab, size=12.5, color=DARK, align=PP_ALIGN.CENTER)
box(s, 7.65, 2.0, 5.08, 4.5, fill=GREEN, radius=0.18)
text(s, 8.05, 2.3, 4.4, 0.3, "OUR GOAL", size=12, bold=True, color=WHITE)
text(s, 8.05, 2.65, 4.4, 1.0, "No. 1 in Armenia's oil market", size=26, bold=True, color=WHITE, line=0.95)
for i, d in enumerate(["Volume", "Revenue", "Reputation", "Technology"]):
    x = 8.05 + (i % 2) * 2.2
    y = 3.9 + (i // 2) * 0.68
    box(s, x, y, 2.05, 0.52, fill=GREEN_DARK, radius=0.26)
    icon(s, "checkCircle", x + 0.1, y + 0.08, 0.36, bg=None, color="w", ratio=1.0)
    text(s, x + 0.55, y + 0.1, 1.45, 0.35, d, size=14, bold=True, color=WHITE)
icon(s, "trendUp", 8.05, 5.5, 0.55, bg=WHITE, color="g")
text(s, 8.75, 5.5, 3.8, 0.55, "Then: take the model beyond Armenia", size=15, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
footer(s, n)
notes(s, """
KAD Motors distributes Castrol lubricants in Armenia. We have field sales reps, a warehouse, delivery, accounting and management, and customers across the regions of Armenia: workshops and garages, oil-change points, auto-parts shops, fleets and car dealers.
Our goal is to become No. 1 in Armenia's oil market: in volume, in revenue, in reputation and in technology. The platform is how we get there, and once the model works here we want to take it further.
""")

# ---------------------------------------------------------------- Built by the business
s = new_slide(); n = nxt()
header(s, "Built by a distributor", "Designed by our whole team, not only by developers")
roles = [("person", "Sales reps", "Visits, orders, cash"), ("target", "Sales director", "Plans, approvals, prices"), ("box", "Warehouse", "Stock, packing, delivery"),
         ("wallet", "Accounting", "Invoices, debt, cash ledger"), ("chart", "CEO & owner", "Goals, margins, decisions")]
for i, (ic, title, body) in enumerate(roles):
    x = 0.6 + i * 2.45
    box(s, x, 1.95, 2.28, 2.25, fill=LIGHT)
    icon(s, ic, x + 0.72, 2.15, 0.84)
    text(s, x + 0.1, 3.1, 2.08, 0.4, title, size=15, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.1, 3.5, 2.08, 0.6, body, size=12, color=GREY, align=PP_ALIGN.CENTER)
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x + 0.94), Inches(4.3), Inches(0.4), Inches(0.42))
    a.fill.solid(); a.fill.fore_color.rgb = GREEN_MID; a.line.fill.background()
box(s, 0.6, 4.85, 12.13, 1.0, fill=GREEN, radius=0.14)
text(s, 0.9, 4.85, 11.6, 1.0, [[("One platform for the whole distributor. ", {"bold": True, "size": 19}), ("Every module started from a real daily problem in our business.", {"size": 16})]], color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.1, 12.1, 0.6, [[("How we build: ", {"bold": True, "color": GREEN_DARK}), ("the team reports a problem on Monday, the fix is in everyone's phone within days. 750+ improvements since August 2026.", {})]], size=14, color=DARK)
footer(s, n)
notes(s, """
This platform was not ordered from an IT agency. It was designed by our own people: sales reps, the sales director, the warehouse and delivery team, accounting, and management, working every day with our developers.
Each role brought its real daily problems: a rep who loses time typing orders twice, an accountant who chases cash, a director who learns about a problem a month too late. Every module started from one of those problems.
That is why the platform covers the needs of a Castrol distributor end to end, and why it keeps improving: problems reported on Monday are usually fixed on everybody's phone within days. More than 750 improvements have shipped since August 2026.
""")

# ---------------------------------------------------------------- Data -> decision -> action
s = new_slide(); n = nxt()
header(s, "Our principle", "Data is worth nothing until it drives action")
loop = [("database", "Collect", "Visits, GPS, photos, orders, payments, ERP"),
        ("sparkles", "Analyse", "AI and rules spot the problem"),
        ("person", "Find cause & owner", "Why it happened, who is responsible"),
        ("tasks", "Task + deadline", "The solution goes to that person"),
        ("clock", "Control", "Reminders until it is done on time"),
        ("warning", "Escalate or report", "Not solved in plan? Management knows")]
bw, gap = 1.88, 0.17
for i, (ic, title, body) in enumerate(loop):
    x = 0.6 + i * (bw + gap)
    hi = i in (2, 3, 4)
    box(s, x, 2.05, bw, 2.75, fill=GREEN if hi else LIGHT)
    icon(s, ic, x + 0.5, 2.3, 0.85, bg=WHITE if hi else GREEN_TINT)
    text(s, x + 0.15, 3.3, bw - 0.3, 0.6, title, size=15, bold=True, color=WHITE if hi else DARK, align=PP_ALIGN.CENTER, line=0.95)
    text(s, x + 0.15, 3.9, bw - 0.3, 0.8, body, size=11.5, color=WHITE if hi else GREY, align=PP_ALIGN.CENTER)
    text(s, x, 1.68, bw, 0.3, str(i + 1), size=13, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
box(s, 0.6, 5.15, 12.13, 0.06, fill=GREEN_MID, radius=0)
text(s, 0.6, 5.35, 12.1, 0.5, "Steps 2 to 6 run automatically, led by our AI deputy director, Vaghinak.", size=17, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.95, 12.1, 0.8, "Managers stop searching for problems. They receive the problem, its cause, the person responsible and the deadline, and they only step in when the plan slips.", size=14, color=GREY, align=PP_ALIGN.CENTER)
footer(s, n)
notes(s, """
This is the most important idea of the platform. Collecting data alone adds no value. Value comes from decisions and actions.
So the platform works as a closed loop. Step one, it collects data: visits with GPS, photos, orders, payments and the daily ERP file. Step two, AI and business rules analyse it and spot problems. Step three, they find the cause and the person responsible. Step four, they give that person a task with a deadline. Step five, they control it: reminders until it is done. Step six, if it is not solved within the plan, they escalate to management or report it.
Steps two to six are run by Vaghinak, our AI deputy director. Managers stop searching for problems; they receive them already analysed, assigned and tracked.
""")

# ---------------------------------------------------------------- Real problems grid
s = new_slide(); n = nxt()
header(s, "Real problems we had to solve", "Ten daily problems of a distributor, before the platform")
probs = [("clock", "Reps start late, skip visits"), ("mapWarning", "No proof a visit happened"), ("wallet", "Overdue debt nobody chases"), ("phone", "Orders by phone, wrong prices"), ("tag", "Discounts that kill margin"),
         ("payment", "Cash stuck with reps"), ("box", "Stock-outs and dead stock"), ("eye", "No view of the shelf, fakes"), ("sheet", "Late reports, Excel chaos"), ("history", "Customers quietly stop buying")]
for i, (ic, lab) in enumerate(probs):
    x = 0.6 + (i % 5) * 2.45
    y = 1.95 + (i // 5) * 2.15
    box(s, x, y, 2.28, 1.95, fill=LIGHT)
    icon(s, ic, x + 0.25, y + 0.25, 0.75, bg=RED_TINT, color="r")
    text(s, x + 0.25, y + 1.1, 1.9, 0.75, lab, size=14.5, bold=True, line=0.95)
box(s, 0.6, 6.3, 12.13, 0.5, fill=RED_TINT, radius=0.1)
text(s, 0.6, 6.3, 12.13, 0.5, [[("Result: ", {"bold": True, "color": RED}), ("management learns about problems a month late, when the money is already lost.", {})]], size=14.5, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, n)
notes(s, """
These are ten real problems we lived with, and that every distributor we know lives with.
Reps start late and skip planned visits. Nobody can prove that a visit happened. Overdue debt sits without anyone chasing it. Orders come by phone with wrong prices. Discounts eat the margin. Cash stays with reps for days. Some products run out while others sit as dead stock. Nobody knows what is on the shelf, or how many fakes are sold. Reports arrive late from a messy Excel. And customers quietly stop buying, and we notice months later.
The common result: management learns about a problem when the money is already lost. The next two slides show how the platform solves each one.
""")

# ---------------------------------------------------------------- Problem -> solution (2 slides)
pairs = [
    [("clock", "Reps start late, skip visits", "Vaghinak logs every arrival and visit, reminds, asks a must-answer question, reports weekly"),
     ("mapWarning", "No proof a visit happened", "GPS check-in only within 300 m of the outlet, with photo and visit outcome"),
     ("wallet", "Overdue debt nobody chases", "Debt aged per invoice; a collection task for the next visit; promise dates; broken promise escalated"),
     ("phone", "Orders by phone, wrong prices", "Orders in the app, at central tier prices, from the usual basket, even offline"),
     ("tag", "Discounts that kill margin", "Up to 3% automatic; more needs approval; below cost only the CEO")],
    [("payment", "Cash stuck with reps", "Every payment tracked from rep to accountant and matched with the ERP automatically"),
     ("box", "Stock-outs and dead stock", "Days of stock per product, a ready reorder list, dead-stock decisions with owners"),
     ("eye", "No view of the shelf, fakes", "Shelf audit at every visit: Castrol presence, full range, fakes, parallel imports, competitors"),
     ("sheet", "Late reports, Excel chaos", "One daily ERP file: checked, reported by the bot, synced to the app in minutes"),
     ("history", "Customers quietly stop buying", "Each customer's order rhythm is learned; when it breaks, a call-or-visit task")]]
for part, rows in enumerate(pairs, 1):
    s = new_slide(); n = nxt()
    header(s, f"Problem → solution ({part}/2)", "Each problem has an automatic answer and an owner")
    text(s, 0.6, 1.75, 4.6, 0.3, "PROBLEM", size=11, bold=True, color=RED)
    text(s, 5.95, 1.75, 6.5, 0.3, "WHAT THE PLATFORM DOES", size=11, bold=True, color=GREEN)
    for i, (ic, prob, sol) in enumerate(rows):
        y = 2.1 + i * 0.93
        box(s, 0.6, y, 4.75, 0.8, fill=RED_TINT, radius=0.1)
        icon(s, ic, 0.72, y + 0.12, 0.56, bg=WHITE, color="r")
        text(s, 1.45, y, 3.8, 0.8, prob, size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.43), Inches(y + 0.22), Inches(0.42), Inches(0.36))
        a.fill.solid(); a.fill.fore_color.rgb = GREEN; a.line.fill.background()
        box(s, 5.95, y, 6.78, 0.8, fill=GREEN_TINT, radius=0.1)
        text(s, 6.2, y, 6.4, 0.8, sol, size=13, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, n)
    notes(s, "\n".join(f"{p}: {sol}." for _, p, sol in rows) + "\nEvery answer has an owner and a deadline, and the platform checks that it was done.")

# ---------------------------------------------------------------- The solution (app + bot)
s = new_slide(); n = nxt()
header(s, "The platform", "One app from a link, an AI deputy director, and a CEO bot")
pillars = [("smartphone", "Field Visits app", "Opens from a link on iPhone and Android. No app store. Works offline."),
           ("vaghinak", "Vaghinak, AI deputy director", "Watches the day, assigns tasks, chases deadlines, escalates."),
           ("bot", "CEO Telegram bot", "Checks the daily ERP file and sends management reports.")]
for i, (ic, title, body) in enumerate(pillars):
    y = 1.95 + i * 1.4
    box(s, 0.6, y, 6.3, 1.25, fill=LIGHT)
    icon(s, ic, 0.85, y + 0.2, 0.85)
    text(s, 2.0, y + 0.17, 4.75, 0.45, title, size=18, bold=True)
    text(s, 2.0, y + 0.6, 4.75, 0.6, body, size=13, color=GREY)
box(s, 0.6, 6.2, 6.3, 0.55, fill=GREEN_TINT, radius=0.1)
icon(s, "sheet", 0.78, 6.25, 0.45, bg=None, color="g", ratio=1.0)
text(s, 1.4, 6.2, 5.4, 0.55, "All fed every day from the same ERP workbook (Excel)", size=13.5, bold=True, color=GREEN_DARK, anchor=MSO_ANCHOR.MIDDLE)
phone(s, "customer", 7.35, 1.45, 5.45)
phone(s, "bot", 10.05, 1.45, 5.45)
text(s, 7.35, 6.72, 5.4, 0.25, "Demo data", size=9, color=GREY, align=PP_ALIGN.CENTER)
footer(s, n)
notes(s, """
The platform has three parts.
The Field Visits app: a progressive web app. A rep opens one link on any iPhone or Android phone. No app store, and it works without signal.
Vaghinak, our AI deputy director: he watches the working day, gives tasks, chases deadlines and escalates when something slips.
The CEO bot in Telegram: every day accounting drops the ERP workbook into the bot; it checks the file, builds the management reports and feeds the app.
""" + DEMO_NOTE)

# ---------------------------------------------------------------- How it works
s = new_slide(); n = nxt()
header(s, "How it works", "One daily file feeds the whole team")
steps = [("sheet", "ERP workbook", "Excel, updated daily by accounting"),
         ("bot", "CEO bot", "Checks the file, builds the reports"),
         ("repeat", "ERP sync", "Debt, sales, prices, stock and plans"),
         ("smartphone", "Field Visits app", "Reps, managers, warehouse, accounting")]
bw, gap = 2.6, 0.57
for i, (ic, title, body) in enumerate(steps):
    x = 0.6 + i * (bw + gap)
    last = i == 3
    box(s, x, 2.1, bw, 2.3, fill=GREEN if last else WHITE, line=None if last else GREEN, line_w=1.5)
    icon(s, ic, x + 0.25, 2.33, 0.75, bg=WHITE if last else GREEN_TINT, color="g")
    text(s, x + 1.12, 2.5, 1.4, 0.4, f"STEP {i + 1}", size=11, bold=True, color=WHITE if last else GREEN)
    text(s, x + 0.25, 3.2, bw - 0.4, 0.45, title, size=18, bold=True, color=WHITE if last else DARK)
    text(s, x + 0.25, 3.62, bw - 0.4, 0.7, body, size=12.5, color=WHITE if last else GREY)
    if i < 3:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.1), Inches(3.05), Inches(gap - 0.2), Inches(0.42))
        a.fill.solid(); a.fill.fore_color.rgb = GREEN; a.line.fill.background()
x2 = 0.6 + 1 * (bw + gap)
a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x2 + bw / 2 - 0.2), Inches(4.5), Inches(0.4), Inches(0.45))
a.fill.solid(); a.fill.fore_color.rgb = GREEN_MID; a.line.fill.background()
box(s, x2 - 0.35, 5.05, bw + 0.7, 0.85, fill=GREEN_TINT, radius=0.12)
icon(s, "send", x2 - 0.2, 5.19, 0.56, bg=WHITE, color="g")
text(s, x2 + 0.5, 5.05, bw - 0.25, 0.85, "Daily reports to management on Telegram", size=13, bold=True, color=GREEN_DARK, anchor=MSO_ANCHOR.MIDDLE)
x4 = 0.6 + 3 * (bw + gap)
a = s.shapes.add_shape(MSO_SHAPE.UP_DOWN_ARROW, Inches(x4 + bw / 2 - 0.2), Inches(4.5), Inches(0.4), Inches(0.45))
a.fill.solid(); a.fill.fore_color.rgb = GREEN_MID; a.line.fill.background()
box(s, x4 - 1.9, 5.05, bw + 1.9, 1.35, fill=LIGHT, radius=0.12)
text(s, x4 - 1.7, 5.15, 4.2, 0.3, "CAPTURED IN THE APP", size=11, bold=True, color=GREY)
for c, xx, yy, ww in zip(["GPS visits", "Photos & shelf audit", "Orders", "Payments"], [x4 - 1.7, x4 - 0.15, x4 - 1.7, x4 - 0.15], [5.5, 5.5, 5.93, 5.93], [1.45, 2.1, 1.45, 2.1]):
    box(s, xx, yy, ww, 0.36, fill=WHITE, line=LINE, radius=0.18, line_w=0.75)
    text(s, xx, yy, ww, 0.36, c, size=11.5, bold=True, color=DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.55, 6.4, 0.35, "Matched by ERP ID, never by name. Excel stays the source of truth.", size=11.5, color=GREY)
footer(s, n)
notes(s, """
The data flow on one page. Accounting keeps the ERP workbook in Excel and uploads it once a day. The CEO bot checks the file and builds the management reports. The same run syncs debt, sales, prices, stock and plans into the app, where every role works.
The app adds what Excel never had: GPS-verified visits, photos and shelf audits, orders and payments captured at the outlet. Customers and products are matched by ERP ID, never by name, and Excel remains the source of truth for debt and sales.
""")

# ---------------------------------------------------------------- Field execution
s = new_slide(); n = nxt()
header(s, "Field execution", "Every visit is verified: GPS check-in, photo and shelf audit")
for name, x, title, body in [("checkin", 1.05, "GPS check-in + photo", "Accepted only within 300 m of the outlet."),
                             ("shelf_audit", 5.2, "Shelf audit at every visit", "Castrol available, full range, fake or parallel import."),
                             ("map", 9.25, "Customer map & CRM", "Every outlet with visit status, debt and plan.")]:
    if name == "shelf_audit":
        w, h = pic(s, os.path.join(SHOTS, "shelf_audit.png"), x + 0.05, 1.95, h=3.85)
    else:
        w, h = phone(s, name, x, 1.75, 4.25)
    cx = x + w / 2
    text(s, cx - 1.75, 6.0, 3.5, 0.35, title, size=16, bold=True, align=PP_ALIGN.CENTER)
    text(s, cx - 1.75, 6.35, 3.5, 0.55, body, size=12.5, color=GREY, align=PP_ALIGN.CENTER)
footer(s, n)
notes(s, """
A rep can only check in when the phone is near the outlet (300 metres by default, set by management), attaches a photo and records the outcome.
At every visit the rep records the shelf: Castrol available or not, full range, fake product, parallel imports (from the USA, Dubai or Russia) and products without the Armenian QR/excise mark, plus competitor brands. Nobody else in our market collects this.
The map is also the CRM: every outlet with its visit status, debt and next planned visit; an automatic planner builds balanced weekly routes for each rep.
""" + DEMO_NOTE)

# ---------------------------------------------------------------- Commercial control
s = new_slide(); n = nxt()
header(s, "Commercial control", "Orders, prices, credit and cash in one controlled flow")
phone(s, "order", 0.55, 1.75, 5.0)
phone(s, "debt", 3.2, 1.75, 5.0)
for i, (ic, title, body) in enumerate([("cart", "Orders in a few taps", "Usual basket and repeat order, even offline."),
                                       ("tag", "Central price lists", "Bronze, Silver and Gold tiers; no typed prices."),
                                       ("checkCircle", "Discount control", "Above 3% needs approval; below cost only the CEO."),
                                       ("wallet", "Credit and cash", "Credit limits, overdue alerts, cash tracked to the accountant."),
                                       ("target", "Targets and bonuses", "Rep targets and points for visits, orders and collections.")]):
    y = 1.95 + i * 0.95
    icon(s, ic, 6.35, y, 0.7)
    text(s, 7.3, y, 5.4, 0.35, title, size=17, bold=True)
    text(s, 7.3, y + 0.36, 5.4, 0.4, body, size=13, color=GREY)
footer(s, n)
notes(s, """
Orders: built in a few taps from the customer's usual basket or a repeat of the last order, even offline. Prices come from central tier price lists; reps never type prices. Discounts up to 3% are automatic; more needs management approval, and a price below our landing cost needs the CEO.
Credit: limits per customer, overdue alerts and collection tasks; every cash payment is tracked from the rep to the accountant. Reps see their own targets and earn points for visits, orders and collections.
""" + DEMO_NOTE)

# ---------------------------------------------------------------- Vaghinak 1: what he controls
s = new_slide(); n = nxt()
header(s, "Vaghinak, AI deputy director", "Vaghinak controls the working day of every person")
phone(s, "vaghinak", 0.55, 1.75, 5.0)
phone(s, "vaghinak_rep", 3.2, 1.75, 5.0)
ctrl = [("clock", "Arrival & lateness", "Logs every start; reports the week"),
        ("mapWarning", "No visit yet", "Asks why at 60 min; must answer at 2 h"),
        ("store", "Too long in the office", "Reminder, then a must-answer question"),
        ("route", "Stopped mid-day", "Visits left? Asks why, keeps the plan"),
        ("mapPinCheck", "Skipped nearby shops", "Planned shop close by? Sends the rep back"),
        ("wallet", "Debt collection", "Task per overdue invoice; promise dates"),
        ("phone", "Calls & reorders", "Customer stopped ordering? Call or visit"),
        ("payment", "Down payments", "Chases unpaid down payments after delivery"),
        ("cart", "Orders waiting", "Too slow? Sales director, then the CEO"),
        ("sheet", "Data & Excel", "Missing data or late ERP file: a task")]
for i, (ic, title, body) in enumerate(ctrl):
    x = 6.3 + (i % 2) * 3.25
    y = 1.85 + (i // 2) * 0.97
    box(s, x, y, 3.12, 0.86, fill=LIGHT, radius=0.1)
    icon(s, ic, x + 0.1, y + 0.15, 0.56)
    text(s, x + 0.78, y + 0.1, 2.3, 0.3, title, size=13, bold=True)
    text(s, x + 0.78, y + 0.42, 2.3, 0.42, body, size=10.5, color=GREY)
text(s, 0.55, 6.72, 5.3, 0.25, "Left: the brief to management. Right: a rep's must-answer question (demo, translated).", size=9, color=GREY)
footer(s, n)
notes(s, """
Meet Vaghinak, our AI operations deputy director. He lives inside the app, with his own page and a badge on the home screen.
He controls the working day of every person. He logs every arrival and every late start, staying quiet within a 30-minute grace and reporting the week. If a rep has not made a first visit after an hour he asks why; after two hours the question becomes mandatory and blocks the app until answered. He notices a rep sitting in the office with visits planned, or stopping in the middle of the day, or leaving planned shops a few hundred metres behind.
He also owns the money and customer follow-ups: a collection task for every overdue invoice, with promise dates; call-or-visit tasks for customers who stopped ordering; chasing down payments after delivery; orders that wait too long go to the sales director and then to the CEO; and missing data or a late ERP file becomes a task for the right person.
On the left is his evening brief to management; on the right, a mandatory question to a rep. In real life he speaks Armenian; these demo messages are translated.
""")

# ---------------------------------------------------------------- Vaghinak 2: closed loop example
s = new_slide(); n = nxt()
header(s, "Vaghinak closes the loop", "From problem to result, without a manager chasing anyone")
flow = [("sparkles", "Detect", "An invoice is 56 days past due"),
        ("person", "Cause & owner", "Oldest unpaid invoice; customer of Rep C"),
        ("tasks", "Task + deadline", "Collect at the next visit, due Thursday"),
        ("clock", "Control", "Rep records: customer promised to pay by the 27th"),
        ("warning", "Escalate", "Promise broken: task reopened, sales director told"),
        ("checkCircle", "Close", "Payment arrives in the ERP: task closes itself")]
bw, gap = 1.88, 0.17
for i, (ic, title, body) in enumerate(flow):
    x = 0.6 + i * (bw + gap)
    warn = i == 4
    done = i == 5
    box(s, x, 1.95, bw, 2.6, fill=RED_TINT if warn else (GREEN if done else LIGHT))
    icon(s, ic, x + 0.5, 2.15, 0.85, bg=WHITE if (warn or done) else GREEN_TINT, color="r" if warn else "g")
    text(s, x + 0.12, 3.12, bw - 0.24, 0.4, title, size=15, bold=True, color=WHITE if done else DARK, align=PP_ALIGN.CENTER)
    text(s, x + 0.12, 3.52, bw - 0.24, 0.95, body, size=11.5, color=WHITE if done else GREY, align=PP_ALIGN.CENTER)
text(s, 0.6, 4.62, 12.1, 0.3, "Example with demo data", size=10, color=GREY)
feats = [("team", "Speaks to each person in their own tone: formal with management, warm and friendly with reps"),
         ("checkCircle", "Staff rate every message (helpful / not relevant), so he keeps improving"),
         ("calendar", "Monday digest: each rep's week, and for management who needs a talk"),
         ("shield", "Starts in test mode; management tunes every threshold before he goes live")]
for i, (ic, lab) in enumerate(feats):
    x = 0.6 + (i % 2) * 6.15
    y = 5.0 + (i // 2) * 0.88
    icon(s, ic, x, y, 0.6)
    text(s, x + 0.75, y, 5.25, 0.62, lab, size=12.5, anchor=MSO_ANCHOR.MIDDLE)
footer(s, n)
notes(s, """
Here is one loop from start to finish, with demo data.
Vaghinak detects that an invoice is 56 days past due. He finds the cause and the owner: it is the oldest unpaid invoice of a customer assigned to Rep C. He creates a task to collect at the next visit, with a deadline. He controls it: the rep records that the customer promised to pay by the 27th. When the promise is broken, the task is reopened and the sales director is informed automatically. When the payment arrives in the ERP, even by bank transfer, the task closes itself.
No manager had to chase anyone. The same loop runs for lateness, visits, reorders, down payments, waiting orders and data quality.
He speaks to each person in their own tone, formal with management and warm with reps; staff can rate each message; every Monday he sends each rep their week and tells management who needs a talk; and he starts in test mode so management can tune every threshold before he goes live.
""")

# ---------------------------------------------------------------- AI & automation
s = new_slide(); n = nxt()
header(s, "AI and automation", "AI and automatic calculations replace hours of manual work")
box(s, 0.6, 1.95, 5.2, 4.8, fill=GREEN, radius=0.16)
icon(s, "sparkles", 0.95, 2.25, 0.85, bg=WHITE)
text(s, 2.0, 2.3, 3.6, 0.8, "KAD AI: ask your data", size=21, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.95, 3.35, 4.6, 1.0, "Management asks in plain language and gets an answer with numbers, in seconds:", size=13.5, color=WHITE)
box(s, 0.95, 4.25, 4.5, 0.9, fill=GREEN_DARK, radius=0.1)
text(s, 1.15, 4.25, 4.1, 0.9, "“Which customers bought less EDGE this quarter than last, and who is their rep?”", size=12.5, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.95, 5.35, 4.6, 1.3, ["Read-only access to the data", "Customers and staff anonymised as codes", "Monthly cost cap, every question logged"], size=12, color=WHITE, spacing=4)
text(s, 6.2, 1.95, 6.5, 0.3, "CALCULATED AUTOMATICALLY", size=12, bold=True, color=GREY)
calcs = ["Landing cost per product: freight, customs, excise, VAT, FX",
         "True cost at the selling price, for every discount decision",
         "Debt ageing per invoice and days past due per customer",
         "Suggested credit limit from each customer's purchases",
         "Each customer's order rhythm and usual basket",
         "Days of stock and a ready reorder list",
         "Rep hours, visits, conversion and weekly scorecard",
         "Fuel allowance from the real GPS route",
         "ERP orders and payments matched to app records",
         "FX moves 2%? A price-review task for the owners"]
for i, c in enumerate(calcs):
    y = 2.35 + i * 0.44
    icon(s, "checkCircle", 6.2, y + 0.03, 0.3, bg=None, color="g", ratio=1.0)
    text(s, 6.65, y, 6.1, 0.4, c, size=12.5)
footer(s, n)
notes(s, """
AI and automation do the work that used to take hours of Excel.
KAD AI lets management ask questions in plain language, for example which customers bought less EDGE this quarter and who their rep is, and get an answer with numbers in seconds. The AI has read-only access; customers and staff are anonymised as codes, so no names leave our server; every question is logged and the monthly cost is capped.
On the right, what the platform calculates automatically: landing cost per product from freight, customs, excise, VAT and the exchange rate; the true cost at the selling price for every discount decision; debt ageing per invoice; suggested credit limits; each customer's order rhythm and usual basket; days of stock and reorder lists; rep hours, visits, conversion and weekly scorecards; fuel allowance from the real GPS route; automatic matching of ERP orders and payments; and a price-review task when the exchange rate moves 2%.
""")

# ---------------------------------------------------------------- Reports
s = new_slide(); n = nxt()
header(s, "Financial and management reports", "Every report a distributor needs, live and on the phone")
phone(s, "reports", 0.55, 1.75, 5.0)
groups = [("wallet", "Money & debt", ["Customer debt, as of any date", "Unpaid invoices by age", "Debt statement PDF for the customer", "Payments and cash custody", "Bank statement matching"]),
          ("chart", "Sales & margin", ["Sales vs plan, by rep and channel", "Litres by product line and region", "Margin and price realisation", "Below-cost sales", "Brand volume"]),
          ("team", "Team & field", ["Weekly scorecard per rep", "Progress by reps, day by day", "Visits and shelf presence", "New customer pipeline", "Fuel allowance"]),
          ("box", "Operations", ["Order flow vs time limits", "Delivery speed", "Stock days and dead stock", "Data quality score", "Daily, weekly, monthly bot briefs"])]
for i, (ic, title, items) in enumerate(groups):
    x = 3.3 + (i % 2) * 4.75
    y = 1.85 + (i // 2) * 2.45
    box(s, x, y, 4.55, 2.3, fill=LIGHT)
    icon(s, ic, x + 0.2, y + 0.18, 0.55)
    text(s, x + 0.9, y + 0.25, 3.5, 0.4, title, size=15, bold=True)
    text(s, x + 0.25, y + 0.85, 4.2, 1.45, items, size=11.5, color=DARK, spacing=1.5)
footer(s, n)
notes(s, """
The platform replaces the monthly Excel report pack with live reports on the phone.
Money and debt: customer debt as of any date, unpaid invoices by age, a one-page debt statement PDF to hand to the customer, payments and cash custody, and bank statement matching.
Sales and margin: sales against plan by rep and channel, litres by product line and region, margin and price realisation, sales below cost, brand volume.
Team and field: the weekly scorecard per rep, progress day by day, visits and shelf presence, the new-customer pipeline and the fuel allowance.
Operations: order flow against time limits, delivery speed, stock days and dead stock, a data-quality score, and the bot's daily, weekly and monthly briefs.
Every report respects roles: a rep sees only their own customers.
""")

# ---------------------------------------------------------------- Management visibility
s = new_slide(); n = nxt()
header(s, "Management visibility", "Management sees every rep, region and KPI, every day")
phone(s, "dashboard", 0.55, 1.75, 5.0)
phone(s, "team", 3.2, 1.75, 5.0)
for i, (ic, lab) in enumerate([("chart", "Sales vs plan"), ("wallet", "Cash collected"), ("mapPinCheck", "Visits & coverage"),
                               ("layers", "Litres by product line"), ("history", "Debt ageing"), ("team", "Rep scorecards")]):
    x = 6.35 + (i % 2) * 3.2
    y = 1.95 + (i // 2) * 1.0
    box(s, x, y, 3.0, 0.82, fill=LIGHT, radius=0.12)
    icon(s, ic, x + 0.14, y + 0.13, 0.56)
    text(s, x + 0.85, y, 2.1, 0.82, lab, size=14.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
box(s, 6.35, 5.1, 6.2, 1.4, fill=GREEN, radius=0.14)
icon(s, "bot", 6.6, 5.37, 0.85, bg=WHITE)
text(s, 7.7, 5.1, 4.7, 1.4, [[("CEO bot on Telegram", {"bold": True, "size": 17})], [("Daily, weekly and monthly reports as soon as the ERP file arrives.", {"size": 13.5})]], color=WHITE, anchor=MSO_ANCHOR.MIDDLE, spacing=3)
footer(s, n)
notes(s, """
The company dashboard and team performance show sales against plan per channel and per rep, cash collected, and how much each rep still needs per working day. Management follows sales versus plan, cash, visits and coverage, litres by product line, debt ageing and rep scorecards, and the CEO bot sends daily, weekly and monthly summaries as soon as the ERP file arrives.
""" + DEMO_NOTE)

# ---------------------------------------------------------------- Unique features
s = new_slide(); n = nxt()
header(s, "What makes it unique", "Twelve things we have not seen in other distributor tools")
uniq = [("vaghinak", "AI deputy director", "Finds problems, assigns, chases, escalates"),
        ("eye", "Shelf & fake tracking", "Fakes, parallel imports, no QR/excise"),
        ("tag", "Price floor", "Below landing cost? Only the CEO approves"),
        ("history", "Debt per invoice", "Real age of every unpaid invoice"),
        ("checkCircle", "Self-closing tasks", "Bank payment in the ERP closes the task"),
        ("mapPinCheck", "Pin review", "GPS evidence finds wrong shop locations"),
        ("route", "Auto route planner", "Balanced weekly routes for every rep"),
        ("truck", "Fuel by real route", "Allowance from the GPS route, city vs highway"),
        ("gauge", "Order speed measured", "Seconds and taps per order, per rep"),
        ("wifioff", "Truly offline", "Visits and orders saved without signal"),
        ("lock", "Emergency lockdown", "One switch locks the app on every phone"),
        ("store", "KAD OS for customers", "Free shop ERP that orders from us (pilot)")]
for i, (ic, title, body) in enumerate(uniq):
    x = 0.6 + (i % 4) * 3.07
    y = 1.85 + (i // 4) * 1.6
    box(s, x, y, 2.92, 1.45, fill=LIGHT)
    icon(s, ic, x + 0.18, y + 0.2, 0.62)
    text(s, x + 0.95, y + 0.2, 1.9, 0.62, title, size=13.5, bold=True, anchor=MSO_ANCHOR.MIDDLE, line=0.95)
    text(s, x + 0.18, y + 0.9, 2.62, 0.5, body, size=11, color=GREY)
footer(s, n)
notes(s, """
Twelve features we believe are unique.
An AI deputy director who finds problems, assigns them, chases and escalates. Shelf tracking that records fakes, parallel imports and products without the Armenian excise mark. A price floor: nothing sells below landing cost without the CEO. Debt aged per invoice, not per customer. Tasks that close themselves when the payment arrives in the ERP, even by bank transfer. Pin review, which uses GPS evidence to find shops whose location is wrong on the map. An automatic route planner. Fuel allowance calculated from the real GPS route, split between city and highway. The speed of every order measured in seconds and taps. A truly offline app. An emergency lockdown that locks the app on every phone with one switch. And KAD OS, a free shop ERP for our own customers, in pilot, which lets them order from us directly and in time gives sell-out data.
""")

# ---------------------------------------------------------------- Illustrative month
s = new_slide(); n = nxt()
header(s, "What management sees", "A month on the platform, at a glance")
badge(s, 8.83, 0.95)
tiles = [("108%", "sales vs plan", "$52.4k of $48.5k"), ("91%", "of sales collected", "$47.7k collected"), ("1,240", "GPS-verified visits", "96% of planned visits"),
         ("64%", "outlets with Castrol on the shelf", "Stage 1 target: 70%"), ("$8.9k", "overdue debt found", "23 collection tasks created"), ("19 of 23", "tasks closed on time", "4 escalated to the sales director")]
for i, (num, lab, sub) in enumerate(tiles):
    x = 0.6 + (i % 3) * 4.1
    y = 1.85 + (i // 3) * 2.2
    box(s, x, y, 3.9, 2.0, fill=GREEN_TINT if i >= 3 else LIGHT)
    text(s, x + 0.3, y + 0.2, 3.4, 0.75, num, size=38, bold=True, color=GREEN)
    text(s, x + 0.3, y + 0.98, 3.4, 0.4, lab, size=14, bold=True)
    text(s, x + 0.3, y + 1.4, 3.4, 0.4, sub, size=12, color=GREY)
text(s, 0.6, 6.35, 12.1, 0.5, "All values are illustrative, in USD, to show what the platform measures. They are not KAD Motors figures.", size=12, bold=True, color=RED)
footer(s, n)
notes(s, """
This slide shows what management sees in a typical month. The values are illustrative only, in US dollars; they are not our company figures.
Top row, the classic numbers: sales against plan, cash collected, GPS-verified visits.
Bottom row, what a traditional report never shows: the share of outlets with Castrol on the shelf against the stage 1 target of 70%; the overdue debt the system found and turned into collection tasks; and how many of those tasks were closed on time, with the rest escalated automatically. That is the data-decision-action loop, measured.
""")

# ---------------------------------------------------------------- 3-stage plan
s = new_slide(); n = nxt()
header(s, "The market plan", "The platform drives a 3-stage plan to win Armenia's shelves")
stages = [("STAGE 1", "Presence", "70% of customers carry at least one Castrol product", "Shelf audit: Castrol available", GREEN_TINT, DARK, GREY),
          ("STAGE 2", "Full range", "Complete Castrol product lines on the shelf", "Shelf audit: full range, plus order mix by line", GREEN_MID, DARK, DARK),
          ("STAGE 3", "Preference", "Customers sell mostly Castrol", "Share of each customer's purchases and sell-out", GREEN, WHITE, WHITE)]
for i, (k, title, target, measure, fill, tc, mc) in enumerate(stages):
    x = 0.6 + i * 2.72
    h = 2.85 + i * 0.75
    y = 6.45 - h
    box(s, x, y, 2.58, h, fill=fill, radius=0.14)
    text(s, x + 0.22, y + 0.22, 2.2, 0.3, k, size=11, bold=True, color=GREEN_DARK if i < 2 else WHITE)
    text(s, x + 0.22, y + 0.5, 2.2, 0.45, title, size=20, bold=True, color=tc)
    text(s, x + 0.22, y + 1.0, 2.15, 1.0, target, size=13.5, bold=True, color=tc)
    text(s, x + 0.22, y + h - 0.85, 2.15, 0.7, [[("Measured by: ", {"bold": True}), (measure, {})]], size=11, color=mc)
w, h = pic(s, os.path.join(SHOTS, "shelf_audit.png"), 9.3, 1.85, h=4.05)
text(s, 9.3 + w / 2 - 1.5, 5.95, 3.0, 0.5, "Recorded at every visit (demo)", size=11.5, color=GREY, align=PP_ALIGN.CENTER)
footer(s, n)
notes(s, """
The platform drives our three-stage market plan. Stage 1, presence: 70% of our customers carry at least one Castrol product, measured outlet by outlet with the shelf audit. Stage 2, full range: complete Castrol product lines on the shelf, measured by the full-range tick and the order mix by product line. Stage 3, preference: customers sell mostly Castrol, measured by the share of each customer's purchases and, through KAD OS, their sell-out.
Vaghinak tracks every stage per outlet, rep and region, and turns gaps into tasks.
""")

# ---------------------------------------------------------------- Roadmap
s = new_slide(); n = nxt()
header(s, "Roadmap", "Next: delivery, warehouse and a full AI management team")
for i, (ic, title, items) in enumerate([("truck", "Driver & delivery", ["Optimised delivery routes", "Arrival detected automatically", "Customer signs on the phone"]),
                                        ("box", "Warehouse", ["Picking and packing in the app", "Live stock behind every order", "Stock counts, even offline"]),
                                        ("sparkles", "AI management team", ["Vaghinak leads expert AI agents", "Sales, finance and stock specialists", "Each agent finds, assigns, controls"])]):
    x = 0.6 + i * 4.1
    box(s, x, 1.95, 3.9, 3.85, fill=LIGHT)
    icon(s, ic, x + 0.3, 2.25, 0.95)
    text(s, x + 0.3, 3.4, 3.4, 0.45, title, size=20, bold=True)
    for j, it in enumerate(items):
        yy = 4.0 + j * 0.55
        icon(s, "checkCircle", x + 0.3, yy + 0.02, 0.3, bg=None, color="g", ratio=1.0)
        text(s, x + 0.75, yy, 3.0, 0.5, it, size=13.5, color=DARK)
box(s, 0.6, 6.05, 12.13, 0.68, fill=GREEN_TINT, radius=0.12)
text(s, 0.9, 6.05, 11.6, 0.68, [[("Today: ", {"bold": True, "color": GREEN_DARK}), ("field sales, orders, cash, reports and Vaghinak.   ", {}), ("Next: ", {"bold": True, "color": GREEN_DARK}), ("a complete operating system for a Castrol distributor.", {})]], size=14.5, anchor=MSO_ANCHOR.MIDDLE)
footer(s, n)
notes(s, """
Next steps. Driver and delivery: optimised routes, automatic arrival detection, and the customer signs for the delivery on the phone. Warehouse: picking and packing in the app, live stock behind every order, stock counts even offline. AI management team: Vaghinak becomes the lead of a team of expert AI agents for sales, finance and stock, each of which finds problems, assigns them and controls the result.
Timing to be confirmed; we ship in small weekly releases.
""")

# ---------------------------------------------------------------- Why + ask
s = new_slide(); n = nxt()
header(s, "Why it matters for Castrol", "Ready to scale, with Castrol shelf data built in")
text(s, 0.6, 1.95, 6, 0.35, "WHAT CASTROL GETS", size=12, bold=True, color=GREY)
for i, (ic, lab) in enumerate([("eye", "Shelf visibility per outlet: presence and range"), ("chart", "Sell-out signals by product line and region"),
                               ("link", "A standard toolkit, opened from a link"), ("shield", "Data stays with the distributor; Castrol sees aggregates")]):
    y = 2.4 + i * 0.95
    icon(s, ic, 0.6, y, 0.72)
    text(s, 1.55, y, 5.0, 0.72, lab, size=15.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
box(s, 6.95, 1.95, 5.78, 4.85, fill=GREEN, radius=0.18)
text(s, 7.35, 2.2, 4.0, 0.3, "OUR PROPOSAL", size=12, bold=True, color=WHITE)
text(s, 7.35, 2.5, 3.2, 0.55, "A pilot partnership", size=25, bold=True, color=WHITE)
logo_spot(s, 10.6, 2.18, 1.9, 0.78, on_dark=True)
for i, a in enumerate(["Monthly Castrol shelf and sell-out dashboard for Armenia", "Pilot with 1–2 more Castrol distributors", "Joint review of results after 90 days"]):
    y = 3.35 + i * 0.82
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.35), Inches(y), Inches(0.5), Inches(0.5))
    c.fill.solid(); c.fill.fore_color.rgb = WHITE; c.line.fill.background()
    text(s, 7.35, y, 0.5, 0.5, str(i + 1), size=15, bold=True, color=GREEN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.05, y - 0.04, 4.45, 0.62, a, size=14.5, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
box(s, 7.35, 5.85, 5.0, 0.68, fill=GREEN_DARK, radius=0.12)
text(s, 7.55, 5.85, 4.7, 0.68, [[("Next step: ", {"bold": True}), ("a 60-minute session with Castrol Global", {})]], size=13.5, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
footer(s, n)
notes(s, """
For Castrol: shelf visibility outlet by outlet; sell-out signals by product line and region, not only sell-in; a standard toolkit that any distributor can start from a link; and a clean data model where operational data stays with the distributor and Castrol receives aggregated dashboards.
Our proposal is a pilot partnership: a monthly Castrol shelf and sell-out dashboard for Armenia, a pilot with one or two more Castrol distributors, and a joint review after 90 days. The next step we ask for is a 60-minute working session with your team.
""")

# ---------------------------------------------------------------- Appendix: architecture
s = new_slide(); n = nxt()
header(s, "Appendix", "Architecture and security")
for i, (ic, title, body) in enumerate([("smartphone", "Progressive web app", "Installs from a link on iPhone, Android and desktop."),
                                       ("wifioff", "Works offline", "Visits and orders wait on the phone and send automatically."),
                                       ("person", "Role-based access", "8 roles; every rule is enforced on the server."),
                                       ("database", "PostgreSQL database", "Runs on the distributor's own cloud server."),
                                       ("lock", "Secure by default", "HTTPS, signed sessions, CSRF protection, rate limits, lockout."),
                                       ("shield", "Data stays with the distributor", "Castrol gets aggregates; AI analysis uses anonymised codes.")]):
    x = 0.6 + (i % 3) * 4.1
    y = 1.95 + (i // 3) * 2.05
    box(s, x, y, 3.9, 1.85, fill=LIGHT)
    icon(s, ic, x + 0.25, y + 0.28, 0.75)
    text(s, x + 1.2, y + 0.3, 2.55, 0.65, title, size=15.5, bold=True, line=0.95)
    text(s, x + 1.2, y + 0.95, 2.55, 0.85, body, size=12, color=GREY)
text(s, 0.6, 6.2, 12.1, 0.5, [[("Quality: ", {"bold": True, "color": DARK}), ("about 1,200 automated tests run on every change; 750+ improvements shipped since August 2026.", {})]], size=12.5, color=GREY)
footer(s, n)
notes(s, """
For digital and IT colleagues: a progressive web app that installs from a link; offline-first, with visits and orders queued on the phone and sent without duplicates; 8 roles with every rule enforced on the server; PostgreSQL on the distributor's own cloud server; HTTPS, signed session cookies, CSRF protection, rate limits and lockout; operational data stays with the distributor and the AI sees anonymised codes only. About 1,200 automated tests run on every change.
""" + DEMO_NOTE)

# Flat look in every renderer: drop the theme style reference (LibreOffice draws its shadow even with an empty effect list).
for sl in prs.slides:
    for el in sl.shapes._spTree.iter():
        if el.tag in (qn("p:sp"), qn("p:cxnSp")):
            st = el.find(qn("p:style"))
            if st is not None:
                el.remove(st)
prs.save(OUT)
print("saved", OUT, len(prs.slides._sldIdLst), "slides")
