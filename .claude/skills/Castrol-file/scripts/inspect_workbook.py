#!/usr/bin/env python3
"""Health check for a Castrol_x.y.z.xlsx workbook (values only, read-only mode, ~20 s).

Usage: python3 inspect_workbook.py path/to/Castrol.xlsx

Prints sheet sizes + header rows (so a renamed/moved column is obvious when the
owner ships a new structure) and runs the data-quality checks listed in
references/data-quality-and-gotchas.md. Needs only openpyxl.
"""
import collections
import sys

from openpyxl import load_workbook

# sheet -> (header row, first data row)
LAYOUT = {
    "Purchase_IN": (2, 3), "Products": (3, 4), "Inventory": (2, 3), "PriceList": (2, 3),
    "Customers": (2, 3), "Debits": (3, 4), "ORDERS": (2, 3), "Cashflow": (9, 10),
}


def rows(ws, header_row, first):
    it = ws.iter_rows(values_only=True)
    allrows = list(it)
    header = [str(h).strip() if h is not None else "" for h in allrows[header_row - 1]]
    out = []
    for r in allrows[first - 1:]:
        if not any(v not in (None, "") for v in r):
            continue
        out.append({header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]})
    return header, out


def main(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    print("SHEETS:", wb.sheetnames)
    data = {}
    for name, (hr, first) in LAYOUT.items():
        if name not in wb.sheetnames:
            print(f"!! MISSING SHEET {name}")
            continue
        header, rs = rows(wb[name], hr, first)
        data[name] = rs
        print(f"\n[{name}] header row {hr}, {len(rs)} data rows\n  columns: {[h for h in header if h]}")

    problems = []
    O, P, Cu, Db, PL = (data.get(k, []) for k in ("ORDERS", "Products", "Customers", "Debits", "PriceList"))

    # 1. Debits 'Assigned To' must be a sales channel, not a tier (VLOOKUP column index drift).
    tiers = {"Gold", "Silver", "Bronze"}
    bad = [r for r in Db if str(r.get("Assigned To") or "") in tiers]
    if bad:
        problems.append(f"Debits 'Assigned To' holds a TIER for {len(bad)} customers (VLOOKUP index should point at Customers 'Assigned To')")

    # 2. ORDERS Sales_Rep must be a channel; tiers there are data-entry slips.
    reps = collections.Counter(str(r.get("Sales_Rep") or "") for r in O)
    slips = {k: v for k, v in reps.items() if k in tiers}
    if slips:
        problems.append(f"ORDERS Sales_Rep contains tier names: {slips}")

    # 3. ProductID mismatches between ORDERS and PriceList (exact vs case-insensitive).
    pl_ids = {str(r.get("ProductID") or "") for r in PL}
    pl_low = {x.lower() for x in pl_ids}
    o_ids = {str(r.get("ProductID") or "") for r in O}
    case_only = sorted(x for x in o_ids if x not in pl_ids and x.lower() in pl_low)
    missing = sorted(x for x in o_ids if x.lower() not in pl_low)
    if case_only:
        problems.append(f"ProductIDs differing only by case (match case-insensitively!): {case_only}")
    if missing:
        problems.append(f"ORDERS ProductIDs absent from PriceList: {missing}")

    # 4. Duplicate ProductIDs in PriceList.
    dups = [k for k, v in collections.Counter(str(r.get("ProductID")) for r in PL).items() if v > 1]
    if dups:
        problems.append(f"Duplicate ProductID rows in PriceList: {dups}")

    # 5. Tier price coverage (0 == empty in PriceList because it is a VLOOKUP).
    def zero(col):
        return sum(1 for r in PL if not (isinstance(r.get(col), (int, float)) and r.get(col) > 0))
    for col in ("Price T1 Bronze", "Price T2 Silver", "Price T3 Gold", "Retail Price"):
        if PL and col in PL[0]:
            problems.append(f"PriceList {col}: {zero(col)}/{len(PL)} empty (0). Bronze empty -> fall back to silver")

    # 6. Error cells in key columns.
    for sheet, cols in (("Debits", ("Debit", "Risk", "Due Days")), ("PriceList", ("Margin", "Discount"))):
        for col in cols:
            n = sum(1 for r in data.get(sheet, []) if isinstance(r.get(col), str) and r.get(col).startswith("#"))
            if n:
                problems.append(f"{sheet}.{col}: {n} Excel error cells")

    # 7. Customers: tier distribution, missing channel/region, duplicates.
    tier = collections.Counter(str(r.get("Tier") or "") for r in Cu)
    problems.append(f"Customers Tier distribution: {dict(tier)}")
    for col in ("Assigned To", "Region", "Subregion", "Phone Number", "TIN"):
        n = sum(1 for r in Cu if not r.get(col))
        problems.append(f"Customers missing {col}: {n}/{len(Cu)}")

    print("\n=== CHECKS ===")
    for p in problems:
        print(" -", p)


if __name__ == "__main__":
    main(sys.argv[1])
