#!/usr/bin/env python3
"""Print, for every table sheet of a Castrol workbook, each column's formula pattern (the dominant
formula with the row number normalised, how many rows use it, and how many rows hold a typed constant
or are empty), plus the parameter cells above the header. Use it to diff formulas between versions.

Usage: python3 dump_formulas.py path/to/Castrol.xlsx [SheetName]     (~60 s, loads the file twice)
Run it from a directory that does not contain other files (Python imports from the cwd).
"""
import collections
import re
import sys

from openpyxl import load_workbook

# sheet -> (header row, first data row)
LAYOUT = {
    "Purchase_IN": (2, 3), "Products": (3, 4), "Inventory": (2, 3), "PriceList": (2, 3),
    "Customers": (2, 3), "Debits": (3, 4), "ORDERS": (2, 3), "Cashflow": (9, 10),
}


def norm(value, row):
    if isinstance(value, str) and value.startswith("="):
        return re.sub(r"(?<![A-Za-z])(\$?[A-Z]{1,3})%d(?!\d)" % row, r"\1{r}", value)
    return "<const>" if value is not None else "<empty>"


def main(path, only=None):
    wb = load_workbook(path)
    wv = load_workbook(path, data_only=True)
    for name, (hdr, first) in LAYOUT.items():
        if only and name != only:
            continue
        ws, vs = wb[name], wv[name]
        last = ws.max_row
        while last > first and all(c.value is None for c in ws[last]):
            last -= 1
        print(f"\n===== {name}: header row {hdr}, data rows {first}-{last}")
        for r in range(1, hdr):
            for c in ws[r]:
                if c.value is not None:
                    text = c.value.text if hasattr(c.value, "text") else c.value
                    print(f"  {c.coordinate}: {str(text)[:160]!r} -> {str(vs[c.coordinate].value)[:30]!r}")
        for c in ws[hdr]:
            if c.value is None:
                continue
            forms, col = collections.Counter(), c.column_letter
            for r in range(first, last + 1):
                forms[norm(ws[f"{col}{r}"].value, r)] += 1
            print(f" {col} [{c.value}] " + " | ".join(f"{n}x {k[:260]}" for k, n in forms.most_common(4)))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
